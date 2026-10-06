#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
工程量汇总统计 Excel 模板生成器
用法:
    python generate_template.py --out 工程量汇总统计_template.xlsx
功能:
    生成带统一表头/样式的空模板，供 aggregate.py 填充或人工填写。
    Sheet 结构：总汇总 / 图纸目录 / 分专业工程量 / 管线长度 / 面积统计
依赖: openpyxl
"""
from __future__ import annotations
import argparse
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[ERROR] 缺少 openpyxl，请执行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# 样式常量
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11, name="微软雅黑")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_FONT = Font(bold=True, size=16, name="微软雅黑", color="1F4E79")
SECTION_FONT = Font(bold=True, size=12, name="微软雅黑")
BODY_FONT = Font(size=10, name="微软雅黑")
BODY_ALIGN = Alignment(vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def _style_header_row(ws, row=1, ncol=6):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER

def _auto_width(ws, min_w=8, max_w=40):
    for col in ws.columns:
        first = col[0]
        col_idx = getattr(first, "column", None)
        if col_idx is None:
            continue
        letter = get_column_letter(col_idx)
        maxlen = 0
        for cell in col:
            v = getattr(cell, "value", None)
            if v is not None:
                s = str(v)
                w = sum(2 if ord(c) > 127 else 1 for c in s)
                maxlen = max(maxlen, w)
        ws.column_dimensions[letter].width = min(max(maxlen + 2, min_w), max_w)

def generate_template(out_path: str):
    wb = Workbook()
    wb.remove(wb.active)

    # --- 总汇总 ---
    ws = wb.create_sheet("总汇总")
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "工程量汇总统计表"
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.append([])
    ws.append(["一、图纸概况"])
    ws.cell(row=ws.max_row, column=1).font = SECTION_FONT
    ws.append(["图纸总数", "", "", "", "", ""])
    ws.append(["涉及专业", "", "", "", "", ""])

    ws.append([])
    ws.append(["二、工程量汇总"])
    ws.cell(row=ws.max_row, column=1).font = SECTION_FONT
    ws.append(["专业", "构件类型数", "构件总数", "管线图层数", "面积图层数", ""])
    _style_header_row(ws, row=ws.max_row, ncol=6)
    ws.append(["建筑", "", "", "", "", ""])
    ws.append(["结构", "", "", "", "", ""])
    ws.append(["给排水", "", "", "", "", ""])
    ws.append(["暖通", "", "", "", "", ""])
    ws.append(["电气", "", "", "", "", ""])
    for row in ws.iter_rows(min_row=ws.max_row - 4, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
    _auto_width(ws)

    # --- 图纸目录 ---
    ws = wb.create_sheet("图纸目录")
    headers = ["序号", "图号", "图名", "专业", "规格/比例", "备注"]
    ws.append(headers)
    _style_header_row(ws, ncol=len(headers))
    for i in range(1, 21):
        ws.append([i, "", "", "", "", ""])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
    _auto_width(ws)

    # --- 分专业工程量（各专业一个 sheet）---
    for disc in ["建筑", "结构", "给排水", "暖通", "电气"]:
        ws = wb.create_sheet(f"{disc}工程量")
        headers = ["序号", "构件名称", "规格", "数量", "单位", "所在图层"]
        ws.append(headers)
        _style_header_row(ws, ncol=len(headers))
        for i in range(1, 51):
            ws.append([i, "", "", "", "个", ""])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = BORDER
                cell.font = BODY_FONT
                cell.alignment = BODY_ALIGN
        _auto_width(ws)

    # --- 管线长度 ---
    ws = wb.create_sheet("管线长度")
    headers = ["序号", "图层", "长度(m)", "单位"]
    ws.append(headers)
    _style_header_row(ws, ncol=len(headers))
    for i in range(1, 31):
        ws.append([i, "", "", "m"])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
    _auto_width(ws)

    # --- 面积统计 ---
    ws = wb.create_sheet("面积统计")
    headers = ["序号", "图层", "面积(m²)", "单位"]
    ws.append(headers)
    _style_header_row(ws, ncol=len(headers))
    for i in range(1, 31):
        ws.append([i, "", "", "m²"])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.font = BODY_FONT
    _auto_width(ws)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path

if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="生成工程量汇总 Excel 模板")
    ap.add_argument("--out", default="工程量汇总统计_template.xlsx", help="输出模板路径")
    args = ap.parse_args()
    path = generate_template(args.out)
    print(f"[OK] 模板已生成: {path}")
    print(f"     Sheet: 总汇总 / 图纸目录 / 各专业工程量 / 管线长度 / 面积统计")
