#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
表格工具箱 —— Excel/CSV 的创建、编辑、多 sheet、样式美化、增删改查
用法:
    # 从零创建
    python doc_table.py create --out 表格.xlsx --title "工程量清单" \\
        --headers "序号,名称,规格,数量,单位" --rows "1,门,M0921,12,樘|2,窗,C1515,24,樘"
    # 读取现有表格
    python doc_table.py read 表格.xlsx [--sheet Sheet1] [--json]
    # 编辑现有表格（增删改查）
    python doc_table.py edit 表格.xlsx --add-row "3,柱,KZ1,8,根"
    python doc_table.py edit 表格.xlsx --update-cell "B2=新名称"
    python doc_table.py edit 表格.xlsx --delete-row 3
    python doc_table.py edit 表格.xlsx --add-sheet "新表" --headers "A,B,C"
    # 从 CSV 转换
    python doc_table.py from-csv 数据.csv --out 表格.xlsx
    # 导出为 CSV
    python doc_table.py to-csv 表格.xlsx --out 数据.csv
功能:
    1. 创建：从零创建带样式的 Excel（表头深蓝白字、隔行变色、边框、冻结首行）
    2. 读取：读取任意 Excel/CSV，输出结构化数据
    3. 编辑：增行、删行、改单元格、加 sheet、改表头
    4. 转换：Excel ↔ CSV
依赖: openpyxl
"""
from __future__ import annotations
import argparse
import csv
import json
import os
import sys
from datetime import datetime
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[ERROR] 缺少依赖 openpyxl，请执行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)
# ---------------------------------------------------------------- 样式常量
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11, name="微软雅黑")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_FONT = Font(size=10, name="微软雅黑")
BODY_ALIGN = Alignment(vertical="center", wrap_text=True)
STRIPE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(bold=True, size=14, name="微软雅黑", color="1F4E79")
# ---------------------------------------------------------------- 工具
def _ensure_dir(path: str):
    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)
def _parse_rows(rows_str: str) -> list:
    """解析 'a,b,c|d,e,f' 格式的行数据。"""
    if not rows_str:
        return []
    out = []
    for line in rows_str.split("|"):
        line = line.strip()
        if line:
            out.append([c.strip() for c in line.split(",")])
    return out
def _auto_width(ws, min_w=8, max_w=40):
    """按内容自动调整列宽。"""
    for col in ws.columns:
        # 合并单元格时 col 元素可能是 MergedCell，用 column 属性兜底
        first = col[0]
        col_idx = getattr(first, "column", None)
        if col_idx is None:
            continue
        letter = get_column_letter(col_idx)
        maxlen = 0
        for cell in col:
            v = getattr(cell, "value", None)
            if v is not None:
                s = str(v)
                w = sum(2 if ord(c) > 127 else 1 for c in s)
                maxlen = max(maxlen, w)
        ws.column_dimensions[letter].width = min(max(maxlen + 2, min_w), max_w)
def _style_sheet(ws, has_title=False):
    """统一美化工作表。"""
    header_row = 2 if has_title else 1
    ncol = ws.max_column
    # 标题行
    if has_title:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        tc = ws.cell(row=1, column=1)
        tc.font = TITLE_FONT
        tc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
    # 表头
    for c in range(1, ncol + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 22
    # 表体
    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            cell.border = BORDER
            if (r - header_row) % 2 == 0:
                cell.fill = STRIPE_FILL
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _auto_width(ws)
# ---------------------------------------------------------------- 创建
def create_table(out_path: str, title: str = None, headers: list = None,
                 rows: list = None, sheet_name: str = "Sheet1") -> str:
    """创建带样式的 Excel 表格。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel sheet 名上限 31 字符
    r = 1
    if title:
        ws.cell(row=1, column=1, value=title)
        r = 2
    if headers:
        for c, h in enumerate(headers, 1):
            ws.cell(row=r, column=c, value=h)
        r += 1
    if rows:
        for row in rows:
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
            r += 1
    _style_sheet(ws, has_title=bool(title))
    _ensure_dir(out_path)
    wb.save(out_path)
    return out_path
# ---------------------------------------------------------------- 读取
def read_table(path: str, sheet: str = None) -> dict:
    """读取 Excel/CSV，返回结构化数据。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            reader = list(csv.reader(f))
        return {"file": path, "sheets": {"CSV": {
            "headers": reader[0] if reader else [],
            "rows": reader[1:] if len(reader) > 1 else [],
            "row_count": max(0, len(reader) - 1),
        }}}
    wb = load_workbook(path, data_only=True)
    result = {"file": path, "sheets": {}}
    targets = [sheet] if sheet else wb.sheetnames
    for sn in targets:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        data = [[getattr(c, "value", None) for c in row] for row in ws.iter_rows()]
        # 去掉全空行
        data = [row for row in data if any(v is not None and str(v).strip() for v in row)]
        headers = []
        rows = []
        if data:
            # 判断第一行是否为标题（只有第一格有值且合并）
            first = data[0]
            non_empty = [v for v in first if v is not None and str(v).strip()]
            if len(non_empty) == 1 and len(first) > 1:
                # data 已是值列表（非 Cell 对象），直接使用
                headers = data[1] if len(data) > 1 else []
                rows = data[2:] if len(data) > 2 else []
            else:
                headers = first
                rows = data[1:]
        result["sheets"][sn] = {
            "headers": [str(h) if h is not None else "" for h in headers],
            "rows": rows,
            "row_count": len(rows),
        }
    return result
# ---------------------------------------------------------------- 编辑
def edit_table(path: str, sheet: str = None, add_row: str = None,
               delete_row: int = None, update_cell: str = None,
               add_sheet: str = None, headers: str = None,
               update_header: str = None, out_path: str = None) -> str:
    """编辑现有 Excel。"""
    if not os.path.isfile(path):
        raise FileNotFoundError(f"文件不存在: {path}")
    wb = load_workbook(path)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    # 新增 sheet
    if add_sheet:
        ws = wb.create_sheet(add_sheet[:31])
        if headers:
            for c, h in enumerate([x.strip() for x in headers.split(",")], 1):
                ws.cell(row=1, column=c, value=h)
        _style_sheet(ws, has_title=False)
    # 新增行
    if add_row:
        vals = [v.strip() for v in add_row.split(",")]
        ws.append(vals)
        _style_sheet(ws, has_title=False)
    # 删除行
    if delete_row:
        ws.delete_rows(delete_row, 1)
        _style_sheet(ws, has_title=False)
    # 更新单元格（格式 B2=新值）
    if update_cell:
        if "=" not in update_cell:
            raise ValueError("--update-cell 格式应为 单元格=值，如 B2=新名称")
        ref, val = update_cell.split("=", 1)
        ws[ref.strip()] = val.strip()
    # 更新表头
    if update_header:
        for c, h in enumerate([x.strip() for x in update_header.split(",")], 1):
            ws.cell(row=1, column=c, value=h)
        _style_sheet(ws, has_title=False)
    out = out_path or path
    _ensure_dir(out)
    wb.save(out)
    return out
# ---------------------------------------------------------------- CSV 互转
def from_csv(csv_path: str, out_path: str, sheet_name: str = "Sheet1") -> str:
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = list(csv.reader(f))
    if not reader:
        raise ValueError("CSV 文件为空")
    headers = reader[0]
    rows = reader[1:]
    return create_table(out_path, headers=headers, rows=rows, sheet_name=sheet_name)
def to_csv(xlsx_path: str, out_path: str, sheet: str = None) -> str:
    data = read_table(xlsx_path, sheet)
    sn = list(data["sheets"].keys())[0]
    info = data["sheets"][sn]
    _ensure_dir(out_path)
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(info["headers"])
        for row in info["rows"]:
            w.writerow(["" if v is None else v for v in row])
    return out_path
# ---------------------------------------------------------------- 入口
def main():
    ap = argparse.ArgumentParser(description="表格工具箱（创建/读取/编辑/转换）")
    sub = ap.add_subparsers(dest="cmd", required=True)
    # create
    p1 = sub.add_parser("create", help="创建新表格")
    p1.add_argument("--out", required=True, help="输出 xlsx 路径")
    p1.add_argument("--title", default=None, help="表格标题")
    p1.add_argument("--headers", default=None, help="表头，逗号分隔")
    p1.add_argument("--rows", default=None, help="数据行，格式 'a,b,c|d,e,f'")
    p1.add_argument("--sheet", default="Sheet1", help="工作表名")
    # read
    p2 = sub.add_parser("read", help="读取表格")
    p2.add_argument("file", help="xlsx/csv 路径")
    p2.add_argument("--sheet", default=None, help="指定工作表")
    p2.add_argument("--json", action="store_true", help="输出 JSON")
    # edit
    p3 = sub.add_parser("edit", help="编辑表格")
    p3.add_argument("file", help="xlsx 路径")
    p3.add_argument("--sheet", default=None, help="指定工作表")
    p3.add_argument("--add-row", default=None, help="新增行，逗号分隔")
    p3.add_argument("--delete-row", type=int, default=None, help="删除行号")
    p3.add_argument("--update-cell", default=None, help="更新单元格，格式 B2=值")
    p3.add_argument("--add-sheet", default=None, help="新增工作表名")
    p3.add_argument("--headers", default=None, help="新表表头，逗号分隔")
    p3.add_argument("--update-header", default=None, help="更新表头，逗号分隔")
    p3.add_argument("--out", default=None, help="另存为（默认覆盖原文件）")
    # from-csv / to-csv
    p4 = sub.add_parser("from-csv", help="CSV 转 Excel")
    p4.add_argument("csv", help="CSV 路径")
    p4.add_argument("--out", required=True, help="输出 xlsx 路径")
    p5 = sub.add_parser("to-csv", help="Excel 转 CSV")
    p5.add_argument("xlsx", help="xlsx 路径")
    p5.add_argument("--out", required=True, help="输出 CSV 路径")
    p5.add_argument("--sheet", default=None, help="指定工作表")
    args = ap.parse_args()
    try:
        if args.cmd == "create":
            headers = [h.strip() for h in args.headers.split(",")] if args.headers else None
            rows = _parse_rows(args.rows) if args.rows else None
            out = create_table(args.out, args.title, headers, rows, args.sheet)
            print(f"[OK] 表格已创建: {out}")
            if headers:
                print(f"     表头: {headers}")
            if rows:
                print(f"     数据行: {len(rows)} 行")
        elif args.cmd == "read":
            data = read_table(args.file, args.sheet)
            if args.json:
                print(json.dumps(data, ensure_ascii=False, indent=2, default=str))
            else:
                for sn, info in data["sheets"].items():
                    print(f"\n===== 工作表: {sn}（{info['row_count']} 行）=====")
                    print("表头:", " | ".join(info["headers"]))
                    for i, row in enumerate(info["rows"][:20], 1):
                        print(f"  {i}.", " | ".join("" if v is None else str(v) for v in row))
                    if info["row_count"] > 20:
                        print(f"  ... 共 {info['row_count']} 行，仅显示前 20 行")
        elif args.cmd == "edit":
            out = edit_table(args.file, args.sheet, args.add_row, args.delete_row,
                             args.update_cell, args.add_sheet, args.headers,
                             args.update_header, args.out)
            print(f"[OK] 表格已更新: {out}")
        elif args.cmd == "from-csv":
            out = from_csv(args.csv, args.out)
            print(f"[OK] 已转换: {out}")
        elif args.cmd == "to-csv":
            out = to_csv(args.xlsx, args.out, args.sheet)
            print(f"[OK] 已转换: {out}")
    except Exception as e:
        print(f"[ERROR] {e}", file=sys.stderr)
        sys.exit(2)
if __name__ == "__main__":
    main()