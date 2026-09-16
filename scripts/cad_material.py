#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
材料识别模块 —— 从图纸文字标注中提取材料信息，生成材料清单
用法:
    python cad_material.py <cio.json> [--dict rules/material_dict.yaml] [--out report_dir] [--excel]
功能:
    1. 加载工程材料词典（正则 + 字段提取）
    2. 扫描图纸全部文字（含图块属性），提取材料信息
    3. 按材料类别汇总，生成材料清单
    4. 支持导出 Excel
设计说明:
    材料信息散落在图纸的文字标注中（如"C30混凝土""DN100镀锌钢管"）。
    通过正则模式匹配 + 命名实体提取，把非结构化文字转为结构化材料清单。
依赖: pyyaml（Excel 导出需 openpyxl）
"""
from __future__ import annotations
import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
try:
    import yaml
except ImportError:
    print("[ERROR] 缺少依赖 pyyaml，请执行: pip install pyyaml", file=sys.stderr)
    sys.exit(1)
# ---------------------------------------------------------------- 材料识别引擎
class MaterialExtractor:
    def __init__(self, dict_path: str):
        with open(dict_path, encoding="utf-8") as f:
            cfg = yaml.safe_load(f)
        self.version = cfg.get("version", "unknown")
        self.patterns = []
        for m in cfg.get("materials", []):
            try:
                rx = re.compile(m["pattern"])
            except re.error as e:
                print(f"[WARN] 正则编译失败 {m.get('category')}: {e}", file=sys.stderr)
                continue
            self.patterns.append({
                "category": m["category"],
                "regex": rx,
                "fields": m.get("fields", {}),
                "raw_pattern": m["pattern"],
            })
    def extract(self, cio: dict) -> dict:
        hits = []
        for e in cio.get("parsed_entities", []):
            attrs = e.get("attributes", {})
            # 收集该实体的所有文本
            texts = []
            if attrs.get("label_text"):
                texts.append(attrs["label_text"])
            for v in (attrs.get("raw_attribs") or {}).values():
                if v:
                    texts.append(str(v))
            for text in texts:
                for p in self.patterns:
                    for m in p["regex"].finditer(text):
                        fields = {}
                        for gname, fname in p["fields"].items():
                            val = m.groupdict().get(gname)
                            if val:
                                fields[fname] = val
                        if not fields:
                            continue
                        hits.append({
                            "entity_id": e["entity_id"],
                            "category": p["category"],
                            "matched_text": m.group(0).strip(),
                            "fields": fields,
                            "source_text": text[:120],
                            "layer": attrs.get("layer_name", ""),
                            "position": e.get("geometry", {}).get("coordinates", [[]])[0],
                        })
        # 按类别汇总
        by_category = defaultdict(list)
        for h in hits:
            by_category[h["category"]].append(h)
        # 去重（同一实体同一匹配文本只记一次）
        summary = {}
        for cat, items in by_category.items():
            seen = set()
            uniq = []
            for it in items:
                key = (it["entity_id"], it["matched_text"])
                if key not in seen:
                    seen.add(key)
                    uniq.append(it)
            # 材料规格统计：按「规格 + 所在实体」去重，避免同一图块的重复属性被多次计数
            spec_counter = Counter()
            seen_spec_entity = set()
            for it in uniq:
                # 同一规格在同一实体上只计一次
                key = (it["matched_text"], it["entity_id"])
                if key in seen_spec_entity:
                    continue
                seen_spec_entity.add(key)
                spec_counter[it["matched_text"]] += 1
            summary[cat] = {
                "count": len(uniq),
                "specs": [{"spec": s, "count": c} for s, c in spec_counter.most_common(50)],
                "items": uniq[:200],
            }
        # 去重后的总命中数（同一实体同一匹配文本只算一次）
        uniq_total = sum(info["count"] for info in summary.values())
        return {
            "dict_version": self.version,
            "total_hits": uniq_total,
            "raw_hits": len(hits),
            "category_count": len(by_category),
            "by_category": summary,
        }
# ---------------------------------------------------------------- Excel 导出
def export_excel(result: dict, out_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    def style(ws, ncol):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        ws.freeze_panes = "A2"
        for r in range(2, ws.max_row + 1):
            for c in range(1, ncol + 1):
                ws.cell(row=r, column=c).border = border
    # Sheet1 材料汇总
    ws1 = wb.active
    ws1.title = "材料汇总"
    ws1.append(["序号", "材料类别", "规格/型号", "出现次数"])
    i = 1
    for cat, info in sorted(result["by_category"].items(), key=lambda x: -x[1]["count"]):
        for s in info["specs"]:
            ws1.append([i, cat, s["spec"], s["count"]])
            i += 1
    style(ws1, 4)
    for col, w in zip("ABCD", [6, 14, 30, 12]):
        ws1.column_dimensions[col].width = w
    # Sheet2 明细
    ws2 = wb.create_sheet("材料明细")
    ws2.append(["序号", "类别", "匹配文本", "提取字段", "所在图层", "源文本"])
    i = 1
    for cat, info in result["by_category"].items():
        for it in info["items"]:
            fields = "、".join(f"{k}={v}" for k, v in it["fields"].items())
            ws2.append([i, cat, it["matched_text"], fields, it["layer"], it["source_text"]])
            i += 1
    style(ws2, 6)
    for col, w in zip("ABCDEF", [6, 12, 22, 26, 18, 40]):
        ws2.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path
# ---------------------------------------------------------------- 报告输出
def write_material_md(result: dict, cio: dict, out_path: str):
    L = []
    A = L.append
    pm = cio.get("project_meta", {})
    A("# CAD 图纸材料识别报告\n")
    A(f"> 文件：`{pm.get('source_file', '')}`　|　专业：{pm.get('discipline', '')}　|　"
      f"分析时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    A("## 一、识别汇总\n")
    A("| 指标 | 值 |")
    A("|---|---:|")
    A(f"| 词典版本 | {result['dict_version']} |")
    A(f"| 材料信息命中数 | {result['total_hits']} |")
    A(f"| 材料类别数 | {result['category_count']} |")
    A("")
    if not result["by_category"]:
        A("## 二、识别结果\n")
        A("未从图纸文字中提取到材料信息。可能原因：图纸文字较少，或材料信息在材料表/图例中（需先解析表格）。\n")
    else:
        A("## 二、材料清单\n")
        for cat, info in sorted(result["by_category"].items(), key=lambda x: -x[1]["count"]):
            uniq_specs = len(info["specs"])
            A(f"### {cat}（{info['count']} 处标注，{uniq_specs} 种规格）\n")
            A("| 规格/型号 | 出现次数 |")
            A("|---|---:|")
            for s in info["specs"][:20]:
                A(f"| {s['spec']} | {s['count']} |")
            A("")
            A(f"> 注：「出现次数」指该规格在图纸中出现的标注次数（同一图块的多处实例分别计数），"
              f"可用于估算构件数量。\n")
        A("## 三、材料明细（前 100 条）\n")
        A("| 类别 | 匹配文本 | 提取字段 | 所在图层 |")
        A("|---|---|---|---|")
        n = 0
        for cat, info in result["by_category"].items():
            for it in info["items"]:
                if n >= 100:
                    break
                fields = "、".join(f"{k}={v}" for k, v in it["fields"].items())
                A(f"| {cat} | {it['matched_text']} | {fields} | {it['layer']} |")
                n += 1
            if n >= 100:
                break
        A("")
    A("## 四、说明\n")
    A("材料信息通过正则模式从图纸文字标注中提取。若图纸含材料表/明细表，"
      "建议先用表格提取功能解析表格，可获得更完整的材料清单。\n")
    A("---\n")
    A("*本报告由 CAD 图纸材料识别模块自动生成。材料词典可在 rules/material_dict.yaml 中扩展。*")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
# ---------------------------------------------------------------- 入口
def main():
    ap = argparse.ArgumentParser(description="CAD 图纸材料识别")
    ap.add_argument("cio", help="CIO JSON 文件路径")
    ap.add_argument("--dict", default=None, help="材料词典 YAML 路径")
    ap.add_argument("--out", default=None, help="输出目录")
    ap.add_argument("--excel", action="store_true", help="导出 Excel 材料清单")
    ap.add_argument("--json", action="store_true", help="输出 JSON")
    args = ap.parse_args()
    dict_path = args.dict
    if not dict_path:
        dict_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                 "rules", "material_dict.yaml")
    if not os.path.isfile(dict_path):
        print(f"[ERROR] 材料词典不存在: {dict_path}", file=sys.stderr)
        sys.exit(2)
    with open(args.cio, encoding="utf-8") as f:
        cio = json.load(f)
    result = MaterialExtractor(dict_path).extract(cio)
    base = os.path.splitext(os.path.basename(args.cio))[0]
    out_dir = args.out or os.path.dirname(os.path.abspath(args.cio))
    os.makedirs(out_dir, exist_ok=True)
    md_path = os.path.join(out_dir, f"{base}_材料识别报告.md")
    write_material_md(result, cio, md_path)
    print(f"[OK] 材料识别报告: {md_path}")
    if args.json:
        json_path = os.path.join(out_dir, f"{base}_材料识别数据.json")
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(f"[OK] JSON 数据: {json_path}")
    if args.excel:
        try:
            xl = os.path.join(out_dir, f"{base}_材料清单.xlsx")
            export_excel(result, xl)
            print(f"[OK] Excel 清单: {xl}")
        except Exception as e:
            print(f"[WARN] Excel 导出失败: {e}", file=sys.stderr)
    print("\n===== 材料识别摘要 =====")
    print(f"命中数    : {result['total_hits']}")
    print(f"材料类别  : {result['category_count']}")
    for cat, info in sorted(result["by_category"].items(), key=lambda x: -x[1]["count"]):
        print(f"   - {cat}: {info['count']} 处")
        for s in info["specs"][:3]:
            print(f"       · {s['spec']} × {s['count']}")
if __name__ == "__main__":
    main()