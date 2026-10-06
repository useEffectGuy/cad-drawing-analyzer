#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
汇总输出脚本 —— 工作流步骤⑦汇总输出
用法:
    python aggregate.py <work目录> --out 工程量汇总统计.xlsx --report 分析报告.md
功能:
    1. 读取 work/ 下的 catalog.json、quantities.json、wbs.json
    2. 生成 Excel：分专业 sheet + 总汇总 sheet（统一表头/样式）
    3. 生成 Markdown 分析报告
依赖: openpyxl
"""
from __future__ import annotations
import argparse
import json
import os
import sys
from datetime import datetime

def _load_json(path: str) -> dict:
    if not os.path.isfile(path):
        return {}
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

def _style_header(ws, ncol):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    header_font = Font(bold=True, color="FFFFFF", size=11, name="微软雅黑")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

def _auto_width(ws, min_w=8, max_w=40):
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        first = col[0]
        col_idx = getattr(first, "column", None)
        if col_idx is None:
            continue
        letter = get_column_letter(col_idx)
        maxlen = 0
        for cell in col:
            v = getattr(cell, "value", None)
            if v is not None:
                s = str(v)
                w = sum(2 if ord(c) > 127 else 1 for c in s)
                maxlen = max(maxlen, w)
        ws.column_dimensions[letter].width = min(max(maxlen + 2, min_w), max_w)

def generate_excel(work_dir: str, out_path: str):
    """生成汇总 Excel。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    catalog = _load_json(os.path.join(work_dir, "catalog.json"))
    quantities = _load_json(os.path.join(work_dir, "quantities.json"))
    wbs = _load_json(os.path.join(work_dir, "wbs.json"))

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body_align = Alignment(vertical="center", wrap_text=True)

    # --- Sheet 1: 图纸目录 ---
    ws = wb.create_sheet("图纸目录")
    headers = ["序号", "图号", "图名", "专业", "规格/比例", "备注"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for i, item in enumerate(catalog.get("items", []), 1):
        ws.append([
            item.get("序号", i),
            item.get("图号", ""),
            item.get("图名", ""),
            item.get("专业", ""),
            item.get("规格", ""),
            item.get("备注", ""),
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = body_align
    _auto_width(ws)

    # --- Sheet 2-N: 分专业工程量 ---
    disc_stats = catalog.get("discipline_stats", {})
    block_table = quantities.get("block_table", [])
    length_table = quantities.get("length_table", [])
    area_table = quantities.get("area_table", [])

    # 按专业分组 block
    from collections import defaultdict
    disc_blocks = defaultdict(list)
    for b in block_table:
        layer = list(b.get("layers", {}).keys())[0] if b.get("layers") else ""
        disc = "其他"
        for d, kws in {
            "建筑": ["墙", "门", "窗", "WALL", "DOOR", "建"],
            "结构": ["柱", "梁", "板", "COLUMN", "BEAM", "结"],
            "给排水": ["水", "PIPE", "DRAIN", "WATER", "排"],
            "电气": ["灯", "插座", "ELEC", "LIGHT", "电"],
            "暖通": ["风", "DUCT", "HVAC", "暖"],
        }.items():
            if any(k.upper() in layer.upper() for k in kws):
                disc = d
                break
        disc_blocks[disc].append(b)

    for disc, blocks in disc_blocks.items():
        ws = wb.create_sheet(f"{disc}工程量"[:31])
        headers = ["序号", "构件名称", "规格", "数量", "单位", "所在图层"]
        ws.append(headers)
        _style_header(ws, len(headers))
        for i, b in enumerate(blocks, 1):
            specs = b.get("specs", {})
            spec_str = "；".join(f"{k}×{v}" for k, v in specs.items() if k != "未标注")
            ws.append([
                i,
                b.get("block_name", ""),
                spec_str,
                b.get("quantity", 0),
                b.get("unit", "个"),
                "、".join(b.get("layers", {}).keys()),
            ])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = body_align
        _auto_width(ws)

    # --- 线长 & 面积 sheet ---
    if length_table:
        ws = wb.create_sheet("管线长度")
        headers = ["序号", "图层", "长度(m)", "单位"]
        ws.append(headers)
        _style_header(ws, len(headers))
        for i, l in enumerate(length_table, 1):
            ws.append([i, l.get("layer", ""), l.get("total_length_m", 0), "m"])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = body_align
        _auto_width(ws)

    if area_table:
        ws = wb.create_sheet("面积统计")
        headers = ["序号", "图层", "面积(m²)", "单位"]
        ws.append(headers)
        _style_header(ws, len(headers))
        for i, a in enumerate(area_table, 1):
            ws.append([i, a.get("layer", ""), a.get("total_area_m2", 0), "m²"])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = body_align
        _auto_width(ws)

    # --- 总汇总 sheet ---
    ws = wb.create_sheet("总汇总", 0)
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = "工程量汇总统计表"
    title_cell.font = Font(bold=True, size=16, name="微软雅黑", color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.append([])
    ws.append(["一、图纸概况"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, name="微软雅黑")
    ws.append(["图纸总数", catalog.get("total_sheets", 0), "", "", "", ""])
    ws.append(["涉及专业", "、".join(disc_stats.keys()), "", "", "", ""])
    for disc, count in disc_stats.items():
        ws.append([f"  {disc}", count, "张", "", "", ""])

    ws.append([])
    ws.append(["二、工程量汇总"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, name="微软雅黑")
    ws.append(["专业", "构件类型数", "构件总数", "管线图层数", "面积图层数", ""])
    _style_header(ws, 6)
    for disc, blocks in disc_blocks.items():
        ws.append([
            disc, len(blocks), sum(b.get("quantity", 0) for b in blocks),
            len(length_table), len(area_table), ""
        ])
    for row in ws.iter_rows(min_row=ws.max_row - len(disc_blocks) + 1, max_row=ws.max_row):
        for cell in row:
            cell.border = border
    _auto_width(ws)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path

def generate_report(work_dir: str, out_path: str) -> str:
    """生成 Markdown 分析报告。"""
    catalog = _load_json(os.path.join(work_dir, "catalog.json"))
    quantities = _load_json(os.path.join(work_dir, "quantities.json"))
    wbs = _load_json(os.path.join(work_dir, "wbs.json"))

    lines = []
    lines.append("# 图纸分析报告")
    lines.append("")
    lines.append(f"**生成时间**：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"**划分标准**：GB 50300-2013《建筑工程施工质量验收统一标准》")
    lines.append("")

    # 一、图纸概况
    lines.append("## 一、图纸概况")
    lines.append("")
    lines.append(f"- **图纸总数**：{catalog.get('total_sheets', 0)} 张")
    lines.append(f"- **目录来源**：{catalog.get('source', '未知')}")
    disc_stats = catalog.get("discipline_stats", {})
    if disc_stats:
        lines.append(f"- **涉及专业**：{len(disc_stats)} 个")
        for disc, count in disc_stats.items():
            lines.append(f"  - {disc}：{count} 张")
    lines.append("")

    # 二、图纸目录
    lines.append("## 二、图纸目录")
    lines.append("")
    lines.append("| 序号 | 图号 | 图名 | 专业 | 规格 | 备注 |")
    lines.append("|---|---|---|---|---|---|")
    for item in catalog.get("items", []):
        lines.append(f"| {item.get('序号','')} | {item.get('图号','')} | {item.get('图名','')} | {item.get('专业','')} | {item.get('规格','')} | {item.get('备注','')} |")
    lines.append("")

    # 三、分专业工程量
    lines.append("## 三、分专业工程量")
    lines.append("")
    s = quantities.get("summary", {})
    lines.append(f"- 构件类型数：{s.get('block_types', 0)}")
    lines.append(f"- 构件总数：{s.get('block_total', 0)}")
    lines.append(f"- 线长图层数：{s.get('length_layers', 0)}")
    lines.append(f"- 面积图层数：{s.get('area_layers', 0)}")
    lines.append("")
    lines.append("### 构件明细（前 20 条）")
    lines.append("")
    lines.append("| 构件名称 | 数量 | 单位 | 规格 |")
    lines.append("|---|---|---|---|")
    for b in quantities.get("block_table", [])[:20]:
        specs = b.get("specs", {})
        spec_str = "；".join(f"{k}×{v}" for k, v in specs.items() if k != "未标注")
        lines.append(f"| {b.get('block_name','')} | {b.get('quantity',0)} | {b.get('unit','个')} | {spec_str} |")
    lines.append("")

    # 四、分部分项划分
    lines.append("## 四、分部分项划分（GB 50300）")
    lines.append("")
    divisions = wbs.get("divisions", {})
    for div, subs in divisions.items():
        lines.append(f"### {div}")
        lines.append("")
        for sub, items in subs.items():
            lines.append(f"- **{sub}**")
            for subitem, entries in items.items():
                lines.append(f"  - {subitem}：{len(entries)} 项")
        lines.append("")

    # 五、结论与建议
    lines.append("## 五、结论与建议")
    lines.append("")
    lines.append(f"本套图纸共 {catalog.get('total_sheets', 0)} 张，涉及 {len(disc_stats)} 个专业，"
                 f"识别构件 {s.get('block_total', 0)} 个，"
                 f"分部分项划分覆盖 {wbs.get('division_count', 0)} 个分部。")
    if catalog.get("out_of_catalog_files"):
        lines.append(f"")
        lines.append(f"⚠️ 发现目录外文件：{catalog['out_of_catalog_files']}，建议核对是否漏登目录。")
    lines.append("")

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return out_path

# ---------------------------------------------------------------- 入口
def main():
    ap = argparse.ArgumentParser(description="汇总输出（工作流步骤⑦）")
    ap.add_argument("work_dir", help="work 目录路径（含 catalog.json/quantities.json/wbs.json）")
    ap.add_argument("--out", required=True, help="输出 Excel 路径")
    ap.add_argument("--report", help="可选：输出 Markdown 报告路径")
    args = ap.parse_args()

    excel_path = generate_excel(args.work_dir, args.out)
    print(f"[OK] Excel 汇总: {excel_path}")

    if args.report:
        report_path = generate_report(args.work_dir, args.report)
        print(f"[OK] 分析报告: {report_path}")

if __name__ == "__main__":
    main()
